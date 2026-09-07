"""
scan_pptx.py — Dump the text content of every slide in a PPTX (read-only).

Used to carry over last month's narrative (projects, status, bullets) into the
new deck.json without opening PowerPoint.

Usage:
  python scan_pptx.py "RG/2026/8. Agosto/9. RG. 25.08.2026 Sistemas 07 - 2026.pptx"
  python scan_pptx.py deck.pptx --slides 2 4 5
"""

import argparse
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(sh.shapes)
        else:
            yield sh


def dump_slide(slide, n):
    print(f"\n{'=' * 60}\nSlide {n}\n{'=' * 60}")
    for sh in iter_shapes(slide.shapes):
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            print(f"  [image] {sh.name}")
            continue
        if getattr(sh, "has_chart", False) and sh.has_chart:
            ch = sh.chart
            try:
                cats = list(ch.plots[0].categories)
                for s in ch.plots[0].series:
                    print(f"  [chart] {s.name}: {dict(zip(cats, s.values))}")
            except Exception:
                print(f"  [chart] {sh.name}")
            continue
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                print("  [table] " + " | ".join(c.text for c in row.cells))
            continue
        if not sh.has_text_frame:
            continue
        lines = [p.text for p in sh.text_frame.paragraphs if p.text.strip()]
        if lines:
            print(f"  [{sh.name}]")
            for line in lines:
                print(f"    {line}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--slides", nargs="*", type=int, help="1-based slide numbers (default: all)")
    args = ap.parse_args()

    prs = Presentation(args.pptx)
    print(f"{args.pptx}  ({len(prs.slides)} slides)")
    wanted = set(args.slides) if args.slides else None
    for i, slide in enumerate(prs.slides, start=1):
        if wanted is None or i in wanted:
            dump_slide(slide, i)


if __name__ == "__main__":
    main()
